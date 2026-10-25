"""Five pitch slides with speaker notes (python-pptx), from real screenshots of the night build.

Run (python-pptx + Pillow on PYTHONPATH), from the repository root:
  python scripts/media/build_slides.py --out <media dir>/AriNav_pitch.pptx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x00, 0x24, 0x78)
ORANGE = RGBColor(0xFF, 0x7A, 0x1A)
INK = RGBColor(0x14, 0x1C, 0x30)
MUTED = RGBColor(0x5C, 0x66, 0x7D)
DOCS = Path("docs/night")

SLIDES = [
    {
        "title": "AriNav: talk to any building",
        "bullets": [
            "Indoor navigation from the building file alone, no beacons",
            "Ask in plain language: 'Where can I get a coffee?'",
            "Official route, guidance, first-person walkthrough",
            "Grounded, accessible, honest about what it knows",
        ],
        "image": "scenes/S1-2-route.png",
        "notes": "Big buildings are hard to navigate, and indoor positioning is expensive. AriNav "
        "needs only the building file. You talk to it like to a colleague at the front desk: it "
        "finds the nearest coffee machine and shows the official route with directions. "
        "Everything it says is grounded in the building data.",
    },
    {
        "title": "Walk it before you go",
        "bullets": [
            "Simulated walkthrough with guidance card and short spoken hints",
            "First person at eye height, minimap, overview at any time",
            "Walk yourself: keyboard or touch joystick, walls stop you",
            "'What's on my right?' answered from pose + geometry",
        ],
        "image": "scenes/S1-4-walking-first-person.png",
        "notes": "You can preview the trip: a clearly labelled simulation walks the route, with "
        "a first-person view and spoken hints. Or walk yourself with keys or a joystick; walls "
        "stop you. Ask what is on your right and the answer comes from the geometry around your "
        "position and view direction, never from guessing what the screen shows.",
    },
    {
        "title": "Accessible by default, honest always",
        "bullets": [
            "Wheelchair: step-free binds every route, never a stairs fallback",
            "Closed lift -> 'no step-free way upstairs' + alternatives",
            "Closures need confirmation; one scenario for map, router, assistant",
            "Outdated 3D model -> live view, clearly marked",
        ],
        "image": "scenes/S2-2-no-step-free-route.png",
        "notes": "Accessibility is a hard rule, not a preference: once someone says they use a "
        "wheelchair, no answer may route them over stairs. When the only lift is closed, AriNav "
        "says so and offers step-free alternatives. Changes to the building need confirmation, and "
        "if the 3D model is outdated, the app shows the live view instead of wrong furniture.",
    },
    {
        "title": "Built on evidence",
        "bullets": [
            "Official kit 1.1.0 router: 655/655 · 1338/1338 · 553/553 · 239/239",
            "140 seeded navigation cases with independent path checks",
            "60 offline state scenarios · 20 live Claude dialogs for $0.30",
            "Browser checks on desktop, 1000x860, phone portrait and landscape",
        ],
        "image": "e2e/HQ-proc-campus-dev-001-1440.png",
        "notes": "We measured instead of claiming. The official deterministic harness is fully "
        "green. A hundred forty seeded navigation cases check every produced path independently. "
        "Sixty offline scenarios test state and guards, twenty live dialogs with Claude Sonnet "
        "cost thirty cents in total, and real browser screenshots cover desktop and phones.",
    },
    {
        "title": "Next: from simulation to the real floor",
        "bullets": [
            "Test phone rotation and microphone on real devices",
            "QR / door anchors as labelled location sources",
            "More buildings through the same deterministic pipeline",
            "Positions stay labelled: simulated, stated or measured",
        ],
        "image": "e2e/M01-portrait-390x844-pov.png",
        "notes": "What is next: real device tests for phone rotation and voice, which we have not "
        "done yet, and simple anchors like QR codes at doors as a labelled location source. The "
        "pipeline already brings new buildings in deterministically. And the rule stays: every "
        "position says where it comes from.",
    },
]


def add_slide(prs: Presentation, spec: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    width, height = prs.slide_width, prs.slide_height
    bar = slide.shapes.add_shape(1, 0, 0, width, Emu(int(height * 0.012)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    image = DOCS / spec["image"]
    left_w = Inches(5.4)
    if image.exists():
        pic = slide.shapes.add_picture(str(image), Inches(6.3), Inches(1.1), width=Inches(6.6))
        if pic.height > Inches(5.6):
            ratio = Inches(5.6) / pic.height
            pic.height = Inches(5.6)
            pic.width = int(pic.width * ratio)
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.9), left_w, Inches(1.6))
    tf = title.text_frame
    tf.word_wrap = True
    tf.text = spec["title"]
    run = tf.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.color.rgb = Pt(36), True, NAVY
    body = slide.shapes.add_textbox(Inches(0.6), Inches(2.7), left_w, Inches(4.2))
    bf = body.text_frame
    bf.word_wrap = True
    for i, text in enumerate(spec["bullets"]):
        para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        para.text = f"•  {text}"
        para.space_after = Pt(12)
        para.runs[0].font.size, para.runs[0].font.color.rgb = Pt(19), INK
    foot = slide.shapes.add_textbox(Inches(0.6), height - Inches(0.6), Inches(8), Inches(0.4))
    foot.text_frame.text = "AriNav · night build · screenshots from the real app"
    foot.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    foot.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
    slide.notes_slide.notes_text_frame.text = spec["notes"]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for spec in SLIDES:
        add_slide(prs, spec)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    print(args.out, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
